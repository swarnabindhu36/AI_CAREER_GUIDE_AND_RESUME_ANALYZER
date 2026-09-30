import os
import pypdf
import docx
from utils.text_utils import clean_text

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

def parse_resume_file(filepath):
    if not os.path.exists(filepath):
        return None, "File does not exist."
    
    ext = filepath.rsplit('.', 1)[-1].lower()
    text = ""
    
    try:
        if ext == 'pdf':
            # Use pdfplumber first if available, fallback to pypdf
            if pdfplumber:
                try:
                    with pdfplumber.open(filepath) as pdf:
                        pages_text = [page.extract_text() or "" for page in pdf.pages]
                        text = "\n".join(pages_text)
                except Exception:
                    text = ""
            
            if not text.strip():
                reader = pypdf.PdfReader(filepath)
                pages_text = [page.extract_text() or "" for page in reader.pages]
                text = "\n".join(pages_text)
                
        elif ext == 'docx':
            doc = docx.Document(filepath)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n".join(paragraphs)
            
        elif ext == 'txt':
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
        else:
            return None, f"Unsupported file format: .{ext}"
            
        cleaned = clean_text(text)
        if not cleaned or len(cleaned) < 30:
            return None, "Could not extract sufficient text from the file. Please ensure it is not scanned/empty."
            
        return cleaned, None
        
    except Exception as e:
        return None, f"Error processing file: {str(e)}"
